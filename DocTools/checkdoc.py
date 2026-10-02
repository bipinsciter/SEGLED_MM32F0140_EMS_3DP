# -*- coding: utf-8 -*-
"""Check the generated document against the firmware facts.

Two things matter: that every identifier the firmware implements appears somewhere in
the document, and that the access and build columns agree with what extract.py found.
"""
import io, json, os, re
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = r"D:\TechnicalData\Yagneshbhai\PRODUCT_CODE\xmegabased\ReleaseCode\SEGLED_MM32F0140_EMS_3DP"
DOC = os.path.join(ROOT, "UART_Protocol.docx")

IDS = json.load(io.open(os.path.join(HERE, 'ids.json'), encoding='utf-8'))

# Read the version from the firmware, not from a literal here - pinning it in the
# checker just moves the staleness from the document into the test.
_const = io.open(os.path.join(ROOT, 'sb_const.h'), encoding='utf-8',
                 errors='surrogateescape').read()


def _fw(name):
    m = re.search(r'^#define\s+' + name + r'\s+(\d+)', _const, re.M)
    return int(m.group(1))


SOFT_VER = _fw('FW_MAJOR') * 100 + _fw('FW_MINOR') * 10 + _fw('FW_PATCH')
doc = Document(DOC)

print("headings")
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name == 'Title':
        print("  " + ("  " * max(0, int(p.style.name[-1]) if p.style.name[-1].isdigit() else 0))
              + p.text)

print()
print("%d tables" % len(doc.tables))

# gather every cell of every table
rows_by_id = {}
for t in doc.tables:
    for r in t.rows:
        cells = [c.text.strip() for c in r.cells]
        if cells and re.match(r'^0x[0-9A-F]{2}$', cells[0]):
            rows_by_id.setdefault(cells[0], []).append(cells)

print("%d identifiers appear in tables" % len(rows_by_id))

fail = 0

# 1. every implemented id is documented
for rec in IDS:
    implemented = rec['read'] or rec['write']
    present = rec['hex'] in rows_by_id
    if implemented and not present:
        print("  MISSING  %s %s is implemented but not in the document"
              % (rec['hex'], rec['name']))
        fail += 1
    if not implemented and not present:
        print("  MISSING  %s %s (unimplemented) not listed either"
              % (rec['hex'], rec['name']))
        fail += 1

# 2. access column agrees, for the rows that carry one
MODE_TEXT = {'DP3': 'DP1+DP2+DP3', 'TEMPRH': 'DP1+Temp+RH', None: 'both'}
checked = 0
for rec in IDS:
    if rec['hex'] not in rows_by_id:
        continue
    for cells in rows_by_id[rec['hex']]:
        if len(cells) != 6:
            continue                      # not a parameter table row
        want_acc = ('RW' if rec['read'] and rec['write']
                    else 'R' if rec['read'] else 'W' if rec['write'] else '-')
        got_acc = cells[3]
        got_mode = cells[4]
        if got_acc != want_acc:
            print("  ACCESS   %s %s: document says %s, code says %s"
                  % (rec['hex'], rec['name'], got_acc, want_acc))
            fail += 1
        if got_mode != MODE_TEXT[rec['mode']]:
            print("  BUILD    %s %s: document says %s, code says %s"
                  % (rec['hex'], rec['name'], got_mode, MODE_TEXT[rec['mode']]))
            fail += 1
        checked += 1

print("checked access and build on %d rows" % checked)

# 3. retired ids must not appear at all
body = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for r in t.rows:
        body += "\n" + " ".join(c.text for c in r.cells)
# Look only where an identifier would be listed - the first cell of a table row.
# The prose carries 0x55 as the checksum seed, which is a different thing entirely.
for gone in ('0x55', '0x56'):
    if gone in rows_by_id:
        print("  STALE    %s is still listed as a parameter" % gone)
        fail += 1
    else:
        print("  ok       %s is not listed as a parameter" % gone)

# 4. spot-check a few facts that must be right
for want, why in [
        (str(SOFT_VER), 'firmware version'),
        ('51 bytes', 'live-value frame length'),
        ('23 bytes', 'serial number frame length'),
        ('XGZP6891D', 'pressure sensor'),
        # the log and feature-word sections were reverse engineered off the wire and
        # are easy to lose in a regeneration; spot-check their load-bearing facts
        ('1507 bytes', 'RAM_ALL bulk reply'),
        ('70 bytes', 'regular-log frame'),
        ('72 bytes', '24 hour ring frame'),
        ('epoch seconds, uint32 little endian', 'the shared 50 byte record'),
        ('0x0002', 'the clock bit of the feature word'),
        ('0x0040', 'the temperature bit of the feature word'),
        ('discarded in silence', 'the customer-password condition on 0x71 and 0x5F'),
        ('AT45DB321D', 'data flash'),
        ('57600', 'default baud')]:
    if want in body:
        print("  ok       %s present (%s)" % (want, why))
    else:
        print("  MISSING  %s (%s)" % (want, why))
        fail += 1

# the sections themselves, not just facts inside them
heads = [p.text for p in doc.paragraphs
         if p.style.name.startswith('Heading') or p.style.name == 'Title']
for want in ('Feature word', 'Reading the logs', 'Which unit a temperature is in'):
    if any(want in h for h in heads):
        print("  ok       section present: %s" % want)
    else:
        print("  MISSING  section: %s" % want)
        fail += 1

print()
print("FAILURES: %d" % fail if fail else "document agrees with the firmware")
