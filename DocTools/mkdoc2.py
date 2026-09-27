# -*- coding: utf-8 -*-
"""Regenerate UART_Protocol.docx from the current firmware.

The parameter tables are driven by ids.json, which extract.py produced by walking the
read and write dispatch switches in main.c - so what the document claims a parameter
does is what the code actually does with it.
"""
import io, json, os, re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = r"D:\TechnicalData\Yagneshbhai\PRODUCT_CODE\xmegabased\ReleaseCode\SEGLED_MM32F0140_EMS_3DP"
OUT = os.path.join(ROOT, "UART_Protocol.docx")

IDS = json.load(io.open(os.path.join(HERE, 'ids.json'), encoding='utf-8'))

# The version is read from the firmware rather than typed here: this document once
# carried a stale one, and a stale version is worse than none.
_const = io.open(os.path.join(ROOT, 'sb_const.h'), encoding='utf-8',
                 errors='surrogateescape').read()
def _fw(name):
    m = re.search(r'^#define\s+' + name + r'\s+(\d+)', _const, re.M)
    return int(m.group(1))
FW = '%d.%d.%d' % (_fw('FW_MAJOR'), _fw('FW_MINOR'), _fw('FW_PATCH'))
SOFT_VER = _fw('FW_MAJOR') * 100 + _fw('FW_MINOR') * 10 + _fw('FW_PATCH')
BYNAME = dict((r['name'], r) for r in IDS)

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x60, 0x60, 0x60)

doc = Document()

st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)

for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.8)
    s.top_margin = s.bottom_margin = Inches(0.7)


def H(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = NAVY
    return h


def P(text='', bold=False, italic=False, size=10, colour=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if colour is not None:
        r.font.color.rgb = colour
    return p


def MONO(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(8)
    return p


def BULLET(text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    return p


def TABLE(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8.5)
            if i == 0:
                run.font.name = 'Consolas'
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ================================================================== title
ti = doc.add_heading('NIYAMA_3DP  RS485 / UART Protocol', level=0)
for r in ti.runs:
    r.font.color.rgb = NAVY

P('Segment-LCD environmental monitor, MM32F0141C6P.   Firmware ' + FW + '.',
  bold=True, size=11)
P('This document is generated from the firmware sources. Every parameter listed was '
  'taken from the read and write dispatch tables in main.c, so it describes what the '
  'code does rather than what it was once intended to do.', italic=True, colour=GREY)

H('1.  Link settings', 1)
TABLE(['Setting', 'Value'],
      [['Signalling', 'RS485 half duplex, or plain UART'],
       ['Baud rate', '57600 by default; settable to 9600 or faster (see 0x41)'],
       ['Frame', '8 data bits, no parity, 1 stop bit, no flow control'],
       ['Device address', '1 by default, 0 addresses every device at once'],
       ['Byte order', 'Binary fields are little endian'],
       ['Epoch', 'Seconds since 1 January 1970, in the device\u2019s own local time']],
      [2.0, 4.6])

H('2.  Frame format', 1)
P('A request:')
MONO('  FF   ID   CMD   PID   [ payload ... ]   CRC   FE\n'
     '  0    1    2     3     4 ...             n-2   n-1')
P('A reply:')
MONO('  FD   ID   CMD   STATUS   PID   [ payload ... ]   CRC   FC\n'
     '  0    1    2     3        4     5 ...             n-2   n-1')
P('Fields', bold=True)
TABLE(['Byte', 'Meaning'],
      [['FF / FD', 'Start of a request / start of a reply'],
       ['ID', 'Device address. The reply always carries the device\u2019s own address, '
              'not the one that was asked for.'],
       ['CMD', '0x10 read, 0x11 write, 0x00 read-with-alarm'],
       ['STATUS', 'Reply only. See section 4.'],
       ['PID', 'Parameter identifier. The reply echoes the one from the request \u2014 '
               'check it, see section 11.'],
       ['CRC', 'See section 3'],
       ['FE / FC', 'End of a request / end of a reply']],
      [1.1, 5.5])

P('Addressing by serial number', bold=True)
P('A request may be addressed to a serial number instead of a device address. It opens '
  'with 0xEA followed by the last 8 characters of the serial number, and closes with '
  '0xEB in place of 0xFE. Only the device whose serial number matches will answer.')

H('3.  Checksum', 1)
P('An additive checksum folded into seven bits. It covers every byte from index 1 up to '
  'but not including the checksum itself.')
MONO('total = 0x55;\n'
     'for (i = 1; i < crc_index; i++) total += frame[i];\n'
     'if (total > 0x7F) total &= 0x7F;')
P('Worked example \u2014 read the firmware version from device 1:')
MONO('  request   FF 01 10 34 1A FE\n'
     '            0x55 + 0x01 + 0x10 + 0x34 = 0x9A, over 0x7F, so & 0x7F = 0x1A')

H('4.  Status byte', 1)
P('A bit field. Zero means everything is well. More than one bit can be set at once.')
TABLE(['Bit', 'Name', 'Meaning'],
      [['0x02', 'INVALID_PARA', 'The device does not support that parameter in this build'],
       ['0x04', 'DP1_FAULTY', 'Differential pressure sensor 1 not responding'],
       ['0x08', 'RH_TEMP_FAULTY', 'Temperature / humidity sensor not responding'],
       ['0x10', 'DP2_FAULTY', 'Differential pressure sensor 2 not responding'],
       ['0x20', 'DP3_FAULTY', 'Differential pressure sensor 3 not responding'],
       ['0x40', 'RTC_INVALID', 'The clock has lost integrity; timestamps are not trustworthy']],
      [0.8, 1.7, 4.1])
P('A reply to an unsupported parameter still arrives, with INVALID_PARA set. There is '
  'one exception, described in section 11.')

H('5.  Reading and writing a value', 1)
P('Read', bold=True)
MONO('  request   FF  ID  10  PID  CRC  FE\n'
     '  reply     FD  ID  10  ST  PID  <ASCII digits>  CRC  FC')
P('The value comes back as ASCII decimal with no padding, and a leading minus sign when '
  'it is negative. Its length therefore varies.')
P('Write', bold=True)
MONO('  request   FF  ID  11  PID  <ASCII digits>  CRC  FE\n'
     '  reply     FD  ID  11  ST  PID  <the same digits>  CRC  FC')
P('The reply echoes the request from byte 3 onward. It confirms the frame was accepted, '
  'not that the value was stored \u2014 a value outside the parameter\u2019s permitted '
  'range is discarded silently. Read the parameter back to be sure.', bold=True)

P('Value range', bold=True)
P('Every ASCII value is parsed into a signed 16 bit integer. Anything outside '
  '\u221232768 to 32767 wraps and is stored as something else entirely, so a host must '
  'not send it.')

H('6.  Scaling', 1)
P('Scaling is not uniform. Each parameter\u2019s factor is given in the tables in '
  'section 8; the groups are:')
TABLE(['Group', 'On the wire', 'Example'],
      [['Alarm setpoints', 'Hundredths of the unit (stored as tenths)',
        '5500 = 55.00 Pa, stored as 550'],
       ['Recorded minimum and maximum', 'Hundredths of the unit', '2524 = 25.24 \u00b0C'],
       ['DP reading clamp (0x6E)', 'Tenths of a Pa', '3000 = 300.0 Pa'],
       ['Per-slot DP offset (0x72)', 'Tenths of a Pa', '\u22120025 = \u22122.5 Pa'],
       ['DP zero offset (0x71)', 'Hundredths of a Pa', '+00125 = 1.25 Pa'],
       ['Temperature / humidity calibration', 'Tenths of the unit', '00250 = 25.0'],
       ['Everything else', 'The plain number', '10 = 10 minutes']],
      [2.2, 2.6, 1.8])

H('7.  Timing', 1)
P('The device serves one message at a time and cannot take in the next while it is '
  'replying. A host should leave roughly 80 ms between transactions.')
P('If a reply does not arrive, let the line go quiet for a few hundred milliseconds '
  'before asking again; that is what clears the device\u2019s receiver.')
P('Firmware before 1.0.4 was considerably less forgiving: a failed checksum latched its '
  'receive-pending flag and gated the receive interrupt permanently, so the unit went '
  'deaf for seconds at a time. Back-to-back frames lost 43 of 60 on the bench. In 1.0.4 '
  'the same test loses 2 of 60.', colour=GREY)
